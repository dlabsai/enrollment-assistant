from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime, tzinfo
from typing import Any
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

from babel.core import Locale, UnknownLocaleError
from babel.dates import format_datetime
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_CELL_MAX_LENGTH = 32_767
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="D9EAF7")
_HEADER_FONT = Font(bold=True)
_CELL_ALIGNMENT = Alignment(vertical="top", wrap_text=True)

ExcelScalar = str | int | float | bool | datetime | None


@dataclass(frozen=True)
class ExcelExportCell:
    value: ExcelScalar
    hyperlink: str | None = None
    number_format: str | None = None


def excel_safe_text(value: str | None, *, truncation_suffix: str | None = None) -> str:
    """Return text that fits in one literal-string Excel cell."""
    if value is None:
        return ""

    text = value
    if len(text) <= EXCEL_CELL_MAX_LENGTH:
        return text
    if truncation_suffix is None:
        return text[:EXCEL_CELL_MAX_LENGTH]
    if len(truncation_suffix) >= EXCEL_CELL_MAX_LENGTH:
        raise ValueError("Excel truncation suffix is too long")
    return text[: EXCEL_CELL_MAX_LENGTH - len(truncation_suffix)] + truncation_suffix


@dataclass(frozen=True)
class BrowserDateTimeFormatter:
    time_zone: tzinfo
    locale_name: str
    format_pattern: str

    @classmethod
    def resolve(cls, *, time_zone: str, locale: str | None) -> BrowserDateTimeFormatter:
        try:
            timezone_obj: tzinfo = ZoneInfo(time_zone)
        except ZoneInfoNotFoundError:
            timezone_obj = UTC

        locale_name = _normalize_browser_locale(locale)
        locale_obj = Locale.parse(locale_name)
        format_pattern = "MMM d, h:mm a"
        if "a" not in locale_obj.time_formats["short"].pattern:
            format_pattern = "MMM d, HH:mm"
        return cls(time_zone=timezone_obj, locale_name=locale_name, format_pattern=format_pattern)

    def format(self, value: datetime) -> str:
        if value.tzinfo is None:
            value = value.replace(tzinfo=UTC)
        localized = value.astimezone(self.time_zone)
        return format_datetime(localized, self.format_pattern, locale=self.locale_name)


def _normalize_browser_locale(value: str | None) -> str:
    if value is None:
        return "en_US"
    normalized = value.replace("-", "_").strip()
    if normalized == "":
        return "en_US"
    try:
        return str(Locale.parse(normalized))
    except UnknownLocaleError, ValueError:
        return "en_US"


class ExcelExportWorkbook:
    """Write-only styled workbook that keeps large exports out of application memory."""

    def __init__(
        self, *, sheet_title: str, headers: tuple[str, ...], column_widths: tuple[float, ...]
    ) -> None:
        if len(headers) != len(column_widths):
            raise ValueError("Excel headers and column widths must have the same length")

        self._workbook = Workbook(write_only=True)
        self._worksheet: Any = self._workbook.create_sheet(sheet_title)
        self._headers = headers
        self._row_count = 0
        self._worksheet.freeze_panes = "A2"
        for column_index, width in enumerate(column_widths, start=1):
            self._worksheet.column_dimensions[get_column_letter(column_index)].width = width

        header_cells: list[Cell] = []
        for header in headers:
            cell = WriteOnlyCell(self._worksheet, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = _CELL_ALIGNMENT
            header_cells.append(cell)
        self._worksheet.append(header_cells)

    def append(self, values: tuple[ExcelExportCell | ExcelScalar, ...]) -> None:
        if len(values) != len(self._headers):
            raise ValueError("Excel row does not match the workbook columns")

        cells: list[Cell] = []
        for value in values:
            export_cell = value if isinstance(value, ExcelExportCell) else ExcelExportCell(value)
            cell = WriteOnlyCell(self._worksheet, value=export_cell.value)
            if isinstance(export_cell.value, str):
                # openpyxl infers leading "=" as a formula. Exported user content is
                # always literal text, so override the inferred type without mutating it.
                cell.data_type = "s"
            if export_cell.hyperlink is not None:
                cell.hyperlink = export_cell.hyperlink
                cell.style = "Hyperlink"
            if export_cell.number_format is not None:
                cell.number_format = export_cell.number_format
            cell.alignment = _CELL_ALIGNMENT
            cells.append(cell)
        self._worksheet.append(cells)
        self._row_count += 1

    def save(self, target: Any) -> None:
        last_column = get_column_letter(len(self._headers))
        self._worksheet.auto_filter.ref = f"A1:{last_column}{self._row_count + 1}"
        self._workbook.save(target)
