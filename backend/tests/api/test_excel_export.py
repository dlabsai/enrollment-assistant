from datetime import UTC, datetime
from io import BytesIO

from openpyxl import load_workbook

from app.api.excel_export import ExcelExportWorkbook, ExportDateTimeFormatter


def test_export_date_uses_the_selected_time_zone() -> None:
    formatter = ExportDateTimeFormatter.resolve(time_zone="America/New_York", locale="en-US")

    assert formatter.format_date(datetime(2026, 1, 1, 2, 30, tzinfo=UTC)) == "2025-12-31"


def test_excel_export_workbook_preserves_literal_string_values() -> None:
    values = ("- Markdown bullet", "=1+1")
    workbook = ExcelExportWorkbook(
        sheet_title="Literal text", headers=("Value",), column_widths=(24,)
    )
    for value in values:
        workbook.append((value,))

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    exported_workbook = load_workbook(buffer, data_only=True, read_only=True)
    worksheet = exported_workbook.active
    assert worksheet is not None
    assert list(worksheet.iter_rows(values_only=True)) == [
        ("Value",),
        ("- Markdown bullet",),
        ("=1+1",),
    ]
    exported_workbook.close()


def test_excel_export_workbook_adds_information_sheet_before_data() -> None:
    workbook = ExcelExportWorkbook(sheet_title="Chats", headers=("Chat",), column_widths=(24,))
    workbook.add_information_sheet(
        (
            ("Chats included", 10_000),
            ("Additional matching chats omitted", "Yes"),
            ("Notice", "Additional matching chats were not included."),
        )
    )
    workbook.append(("Example chat",))

    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    exported_workbook = load_workbook(buffer, data_only=True, read_only=True)
    assert exported_workbook.sheetnames == ["Export information", "Chats"]
    assert list(exported_workbook["Export information"].iter_rows(values_only=True)) == [
        ("Item", "Value"),
        ("Chats included", 10_000),
        ("Additional matching chats omitted", "Yes"),
        ("Notice", "Additional matching chats were not included."),
    ]
    assert list(exported_workbook["Chats"].iter_rows(values_only=True)) == [
        ("Chat",),
        ("Example chat",),
    ]
    exported_workbook.close()
